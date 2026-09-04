"""Local, source-labelled material exports for the investment demo."""

from __future__ import annotations

import csv
import io
import uuid
from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Mapping

from .knowledge import KnowledgeRepository


MATERIAL_TEMPLATES: dict[str, dict[str, str]] = {
    "chain": {
        "label": "产业链靶向招商",
        "title": "产业链靶向招商简报",
        "intro": "围绕企业产品、能力和潜在产业链角色，形成可追溯的首轮靶向招商线索。",
        "company_heading": "一、企业与产业链线索",
        "park_heading": "二、园区对标与独资设立洽谈",
        "landing_heading": "三、本地化落地配套核验包",
    },
    "park": {
        "label": "独资设立园区洽谈",
        "title": "独资设立园区洽谈简报",
        "intro": "围绕园区适配、独资设立事项与分阶段洽谈路径组织本地公开资料。",
        "company_heading": "一、可协同企业线索",
        "park_heading": "二、园区对标与独资设立洽谈",
        "landing_heading": "三、项目落地事项核验包",
    },
    "landing": {
        "label": "落地配套协同",
        "title": "项目落地配套协同简报",
        "intro": "围绕政务办事、生活配套和交通区位，形成可分工、可核验的项目落地事项清单。",
        "company_heading": "一、项目关联企业线索",
        "park_heading": "二、园区与项目边界核验",
        "landing_heading": "三、政务、生活与交通配套",
    },
}

class ExportService:
    """Create a concise Word/PPTX brief or a standardised target-list CSV."""

    def __init__(self, repository: KnowledgeRepository) -> None:
        self.repository = repository
        self.output_dir = repository.data_dir / "exports"
        self.output_dir.mkdir(parents=True, exist_ok=True)

    def target_list_csv(self, company_ids: Iterable[str], *, actor: str) -> tuple[Path, str]:
        companies = self.repository.company_by_ids(company_ids)
        if not companies:
            raise ValueError("请至少选择一家已入库企业")
        filename = self._filename("靶向招商清单", "csv")
        path = self.output_dir / filename
        buffer = io.StringIO(newline="")
        writer = csv.DictWriter(
            buffer,
            fieldnames=["企业ID", "企业名称", "区域", "产业赛道", "细分行业", "潜在产业链角色", "产品/能力", "资料可信度", "来源标题", "来源链接", "使用说明"],
        )
        writer.writeheader()
        for company in companies:
            source = (company.get("sources") or [{}])[0]
            writer.writerow(
                {
                    key: self._csv_safe(value)
                    for key, value in {
                    "企业ID": company.get("company_id", ""),
                    "企业名称": company.get("name", ""),
                    "区域": company.get("district", ""),
                    "产业赛道": company.get("sector", ""),
                    "细分行业": company.get("industry", ""),
                    "潜在产业链角色": "、".join(company.get("supply_chain_role", [])),
                    "产品/能力": "、".join([*company.get("products", []), *company.get("capabilities", [])]),
                    "资料可信度": company.get("confidence", ""),
                    "来源标题": source.get("title", ""),
                    "来源链接": source.get("url", ""),
                    "使用说明": "潜在线索，需进一步核验，不代表已建立商业关系。",
                    }.items()
                }
            )
        path.write_text("\ufeff" + buffer.getvalue(), encoding="utf-8")
        export_id = self.repository.record_export(export_type="target_list", fmt="csv", file_path=str(path), actor=actor, record_ids=[item["company_id"] for item in companies])
        return path, export_id

    @staticmethod
    def template_options() -> list[dict[str, str]]:
        return [{"key": key, "label": value["label"], "title": value["title"]} for key, value in MATERIAL_TEMPLATES.items()]

    def material(self, company_ids: Iterable[str], *, project_name: str, fmt: str, actor: str, template_key: str = "chain") -> tuple[Path, str]:
        companies = self.repository.company_by_ids(company_ids)
        if not companies:
            raise ValueError("请至少选择一家已入库企业")
        template = MATERIAL_TEMPLATES.get(template_key)
        if template is None:
            raise ValueError("未知的招商材料模板")
        bundle = self._bundle(companies, project_name, template_key=template_key, template=template)
        if fmt == "docx":
            path = self._write_docx(bundle)
        elif fmt == "pptx":
            path = self._write_pptx(bundle)
        else:
            raise ValueError("仅支持 DOCX 或 PPTX")
        export_id = self.repository.record_export(export_type=f"investment_brief:{template_key}", fmt=fmt, file_path=str(path), actor=actor, record_ids=[item["company_id"] for item in companies])
        return path, export_id

    def _bundle(self, companies: list[dict[str, Any]], project_name: str, *, template_key: str, template: Mapping[str, str]) -> dict[str, Any]:
        sectors = {str(company.get("sector") or "") for company in companies if str(company.get("sector") or "")}
        # A single selected industry can make the exported park matrix more
        # useful without inferring facts for a mixed-sector selection.
        park_sector = next(iter(sectors)) if len(sectors) == 1 else None
        parks = self.repository.compare_parks([], sector=park_sector, project_name=project_name)
        districts = {str(company.get("district") or "") for company in companies if str(company.get("district") or "")}
        # When all selected enterprises are in one district, carry that scope
        # into the exported landing package.  Mixed selections remain honestly
        # labelled as citywide rather than pretending to be a district brief.
        landing_district = next(iter(districts)) if len(districts) == 1 else "常州市"
        landing = self.repository.landing_package(project_name=project_name, district=landing_district)
        sources = self._unique_sources(
            source
            for company in companies
            for source in company.get("sources", [])
        )
        sources.extend(self._unique_sources(source for park in parks.get("parks", []) for source in park.get("sources", [])))
        sources.extend(self._unique_sources(source for section in landing.get("sections", []) for item in section.get("items", []) for source in item.get("sources", [])))
        return {
            "project_name": project_name,
            "created_date": datetime.now().strftime("%Y-%m-%d"),
            "template_key": template_key,
            "template": dict(template),
            "companies": companies,
            "parks": parks,
            "landing": landing,
            "sources": self._unique_sources(sources),
        }

    def _write_docx(self, bundle: Mapping[str, Any]) -> Path:
        try:
            from docx import Document
            from docx.enum.section import WD_SECTION
            from docx.enum.text import WD_ALIGN_PARAGRAPH
            from docx.oxml import OxmlElement
            from docx.oxml.ns import qn
            from docx.shared import Inches, Pt, RGBColor
        except ImportError as exc:
            raise RuntimeError("缺少 python-docx，请按 environment.yaml 安装环境") from exc

        doc = Document()
        section = doc.sections[0]
        section.top_margin = Inches(1)
        section.bottom_margin = Inches(1)
        section.left_margin = Inches(1)
        section.right_margin = Inches(1)
        section.header_distance = Inches(0.492)
        section.footer_distance = Inches(0.492)

        normal = doc.styles["Normal"]
        normal.font.name = "Calibri"
        normal.font.size = Pt(11)
        normal.paragraph_format.space_after = Pt(6)
        normal.paragraph_format.line_spacing = 1.10
        for style_name, size, color, before, after in (
            ("Heading 1", 16, RGBColor(46, 116, 181), 16, 8),
            ("Heading 2", 13, RGBColor(46, 116, 181), 12, 6),
        ):
            style = doc.styles[style_name]
            style.font.name = "Calibri"
            style.font.size = Pt(size)
            style.font.color.rgb = color
            style.paragraph_format.space_before = Pt(before)
            style.paragraph_format.space_after = Pt(after)

        header = section.header.paragraphs[0]
        header.text = "常州产业招商知识平台 · 本地演示"
        header.alignment = WD_ALIGN_PARAGRAPH.RIGHT
        header.runs[0].font.size = Pt(9)
        header.runs[0].font.color.rgb = RGBColor(89, 99, 110)
        footer = section.footer.paragraphs[0]
        footer.text = "公开资料可追溯；潜在匹配不代表已建立商业关系。"
        footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
        footer.runs[0].font.size = Pt(8.5)

        template = bundle["template"]
        title = doc.add_paragraph()
        title.alignment = WD_ALIGN_PARAGRAPH.LEFT
        title.paragraph_format.space_before = Pt(14)
        title.paragraph_format.space_after = Pt(4)
        run = title.add_run(str(template["title"]))
        run.bold = True
        run.font.name = "Calibri"
        run.font.size = Pt(23)
        run.font.color.rgb = RGBColor(11, 37, 69)
        subtitle = doc.add_paragraph()
        subtitle.paragraph_format.space_after = Pt(14)
        subtitle_run = subtitle.add_run(f"项目：{bundle['project_name']}  |  生成日期：{bundle['created_date']}")
        subtitle_run.font.size = Pt(11)
        subtitle_run.font.color.rgb = RGBColor(89, 99, 110)
        self._paragraph_rule(subtitle, "D7DBE2")

        doc.add_heading(str(template["company_heading"]), level=1)
        doc.add_paragraph(str(template["intro"]))
        doc.add_paragraph("以下企业为已入库公开资料的潜在线索。请在正式洽谈前核验产品规格、产能、认证和商务意向。")
        table = doc.add_table(rows=1, cols=4)
        table.style = "Table Grid"
        table.autofit = False
        widths = [Inches(1.55), Inches(1.45), Inches(2.05), Inches(1.45)]
        for index, width in enumerate(widths):
            table.columns[index].width = width
        headers = ["企业", "产业/区域", "公开产品与能力", "来源"]
        for cell, label in zip(table.rows[0].cells, headers):
            cell.text = label
            self._shade_cell(cell, "F2F4F7")
            for run in cell.paragraphs[0].runs:
                run.bold = True
                run.font.size = Pt(9)
        for company in bundle["companies"]:
            cells = table.add_row().cells
            source = (company.get("sources") or [{}])[0]
            values = [
                company.get("name", ""),
                " · ".join(filter(None, [company.get("sector", ""), company.get("district", "")])),
                "、".join([*company.get("products", []), *company.get("capabilities", [])]) or "待补录",
                source.get("source_id", "") or "待核验",
            ]
            for cell, value in zip(cells, values):
                cell.text = str(value)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.font.size = Pt(8.5)
        citation = doc.add_paragraph("企业表来源：详见文末“来源索引”。")
        citation.paragraph_format.space_before = Pt(4)
        citation.paragraph_format.space_after = Pt(4)
        citation.runs[0].font.size = Pt(8.5)
        citation.runs[0].font.color.rgb = RGBColor(89, 99, 110)

        doc.add_heading(str(template["park_heading"]), level=1)
        doc.add_paragraph(bundle["parks"].get("caveat", ""))
        benchmark = bundle["parks"].get("benchmark_model", {})
        benchmark_rows = benchmark.get("rows", []) if isinstance(benchmark, Mapping) else []
        if benchmark_rows:
            doc.add_heading("公开证据对标矩阵", level=2)
            interpretation = str(benchmark.get("score_interpretation", ""))
            if interpretation:
                paragraph = doc.add_paragraph(interpretation)
                paragraph.runs[0].font.size = Pt(8.5)
                paragraph.runs[0].font.color.rgb = RGBColor(89, 99, 110)
            for row in benchmark_rows:
                pending = "、".join(str(item) for item in row.get("pending_dimensions", []) if str(item)) or "无"
                paragraph = doc.add_paragraph(style="List Bullet")
                paragraph.add_run(f"{row.get('park_name', '园区')}：资料完备度 {row.get('evidence_completeness_percent', '—')}%；待核验：{pending}。")
        for point in bundle["parks"].get("negotiation_points", []):
            paragraph = doc.add_paragraph(style="List Bullet")
            paragraph.paragraph_format.space_after = Pt(4)
            paragraph.add_run(point)
        doc.add_heading("建议阶段", level=2)
        for phase in bundle["parks"].get("phased_plan", []):
            paragraph = doc.add_paragraph()
            lead = paragraph.add_run(f"{phase['phase']}：{phase['goal']}。")
            lead.bold = True
            paragraph.add_run("；".join(phase.get("actions", [])))

        doc.add_heading(str(template["landing_heading"]), level=1)
        checklist = bundle["landing"].get("action_checklist", [])
        if checklist:
            doc.add_heading("项目化核验清单", level=2)
            for item in checklist:
                paragraph = doc.add_paragraph(style="List Bullet")
                lead = paragraph.add_run(f"{item.get('phase', '待安排')} · {item.get('workstream', '核验事项')}：")
                lead.bold = True
                paragraph.add_run(f"{item.get('action', '')} 建议对接：{item.get('suggested_owner', '待明确')}。")
        for section_data in bundle["landing"].get("sections", []):
            doc.add_heading(str(section_data.get("category", "配套")), level=2)
            for item in section_data.get("items", []):
                paragraph = doc.add_paragraph()
                lead = paragraph.add_run(f"{item.get('name', '')}：")
                lead.bold = True
                paragraph.add_run(str(item.get("summary", "")))

        doc.add_heading("四、来源索引", level=1)
        for source in bundle["sources"]:
            paragraph = doc.add_paragraph()
            paragraph.paragraph_format.space_after = Pt(4)
            lead = paragraph.add_run(f"[{source.get('source_id', '')}] {source.get('title', '')} ")
            lead.bold = True
            paragraph.add_run(str(source.get("url", "")))

        path = self.output_dir / self._filename("招商研判简报", "docx")
        doc.save(path)
        return path

    def _write_pptx(self, bundle: Mapping[str, Any]) -> Path:
        try:
            from pptx import Presentation
            from pptx.dml.color import RGBColor
            from pptx.enum.text import PP_ALIGN
            from pptx.util import Inches, Pt
        except ImportError as exc:
            raise RuntimeError("缺少 python-pptx，请按 environment.yaml 安装环境") from exc

        presentation = Presentation()
        presentation.slide_width = Inches(13.333)
        presentation.slide_height = Inches(7.5)
        blank = presentation.slide_layouts[6]
        navy = RGBColor(11, 37, 69)
        blue = RGBColor(46, 116, 181)
        muted = RGBColor(89, 99, 110)

        template = bundle["template"]
        slide = presentation.slides.add_slide(blank)
        self._slide_text(slide, 0.75, 1.25, 11.8, 0.7, str(template["title"]), 45, navy, bold=True)
        self._slide_text(slide, 0.75, 2.15, 11.8, 0.45, str(bundle["project_name"]), 26, blue)
        self._slide_text(slide, 0.75, 6.5, 11.8, 0.3, "常州产业招商知识平台 · 仅使用可追溯公开资料", 16, muted)

        slide = presentation.slides.add_slide(blank)
        self._slide_title(slide, str(template["company_heading"]))
        y = 1.65
        for company in bundle["companies"][:5]:
            source = (company.get("sources") or [{}])[0]
            title = f"{company.get('name', '')} · {company.get('district', '区域待核验')}"
            body = "、".join([*company.get("products", []), *company.get("capabilities", [])]) or "公开产品/能力待补录"
            self._slide_text(slide, 0.85, y, 11.6, 0.3, title, 22, navy, bold=True)
            self._slide_text(slide, 1.15, y + 0.35, 11.2, 0.3, f"{body}  |  来源：{source.get('source_id', '待核验')}", 16, muted)
            y += 0.95

        slide = presentation.slides.add_slide(blank)
        self._slide_title(slide, str(template["park_heading"]))
        y = 1.55
        for park in bundle["parks"].get("parks", []):
            metrics = "；".join(f"{item.get('label')}：{item.get('value')}" for item in park.get("metrics", []))
            self._slide_text(slide, 0.85, y, 11.6, 0.3, f"{park.get('name')} · {park.get('district')}", 24, navy, bold=True)
            self._slide_text(slide, 1.15, y + 0.38, 11.1, 0.52, metrics or "公开指标待补录", 16, muted)
            y += 1.25
        benchmark = bundle["parks"].get("benchmark_model", {})
        benchmark_rows = benchmark.get("rows", []) if isinstance(benchmark, Mapping) else []
        if benchmark_rows:
            summary = "；".join(
                f"{row.get('park_name', '园区')} {row.get('evidence_completeness_percent', '—')}%（待核验：{'、'.join(row.get('pending_dimensions', [])) or '无'}）"
                for row in benchmark_rows
            )
            self._slide_text(slide, 0.85, 5.55, 11.4, 0.52, f"公开资料完备度：{summary}", 14, muted)

        slide = presentation.slides.add_slide(blank)
        self._slide_title(slide, str(template["landing_heading"]))
        checklist = bundle["landing"].get("action_checklist", [])
        if checklist:
            first = checklist[0]
            self._slide_text(
                slide,
                0.8,
                1.25,
                11.7,
                0.34,
                f"优先核验：{first.get('workstream', '项目事项')} · {first.get('action', '')}",
                16,
                muted,
            )
        x = 0.8
        for section_data in bundle["landing"].get("sections", []):
            self._slide_text(slide, x, 1.75, 3.7, 0.35, str(section_data.get("category", "配套")), 24, blue, bold=True)
            # A slide is a concise briefing rather than the full landing-service
            # catalogue.  Keeping two entries and a bounded summary prevents a
            # third long public-record description from colliding with the next
            # entry after PowerPoint lays out Chinese text.
            y = 2.25
            for item in section_data.get("items", [])[:2]:
                self._slide_text(slide, x, y, 3.7, 0.3, str(item.get("name", "")), 17, navy, bold=True)
                self._slide_text(slide, x, y + 0.32, 3.7, 0.72, str(item.get("summary", ""))[:48], 14, muted)
                y += 1.45
            x += 4.15

        slide = presentation.slides.add_slide(blank)
        self._slide_title(slide, "来源索引与使用边界")
        self._slide_text(slide, 0.85, 1.55, 11.4, 0.45, "潜在匹配、产业关系与园区适配均需以所列公开来源和后续业务核验为准。", 20, navy)
        y = 2.3
        for source in bundle["sources"][:8]:
            self._slide_text(slide, 0.95, y, 11.1, 0.26, f"[{source.get('source_id', '')}] {source.get('title', '')} — {source.get('url', '')}", 14, muted)
            y += 0.46

        path = self.output_dir / self._filename("招商研判简报", "pptx")
        presentation.save(path)
        return path

    @staticmethod
    def _slide_text(slide: Any, x: float, y: float, width: float, height: float, text: str, size: float, color: Any, *, bold: bool = False) -> None:
        from pptx.enum.text import PP_ALIGN
        from pptx.util import Inches, Pt

        shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
        frame = shape.text_frame
        frame.clear()
        frame.word_wrap = True
        paragraph = frame.paragraphs[0]
        paragraph.alignment = PP_ALIGN.LEFT
        run = paragraph.add_run()
        run.text = text
        run.font.name = "Aptos"
        run.font.size = Pt(size)
        run.font.bold = bold
        run.font.color.rgb = color

    @classmethod
    def _slide_title(cls, slide: Any, text: str) -> None:
        from pptx.dml.color import RGBColor

        cls._slide_text(slide, 0.75, 0.55, 11.85, 0.55, text, 36, RGBColor(11, 37, 69), bold=True)

    @staticmethod
    def _shade_cell(cell: Any, fill: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        properties = cell._tc.get_or_add_tcPr()
        shading = OxmlElement("w:shd")
        shading.set(qn("w:fill"), fill)
        properties.append(shading)

    @staticmethod
    def _paragraph_rule(paragraph: Any, color: str) -> None:
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn

        properties = paragraph._p.get_or_add_pPr()
        borders = OxmlElement("w:pBdr")
        bottom = OxmlElement("w:bottom")
        bottom.set(qn("w:val"), "single")
        bottom.set(qn("w:sz"), "6")
        bottom.set(qn("w:space"), "8")
        bottom.set(qn("w:color"), color)
        borders.append(bottom)
        properties.append(borders)

    @staticmethod
    def _filename(stem: str, suffix: str) -> str:
        timestamp = datetime.now().strftime("%Y%m%d-%H%M%S")
        return f"{stem}-{timestamp}-{uuid.uuid4().hex[:6]}.{suffix}"

    @staticmethod
    def _csv_safe(value: Any) -> str:
        """Prevent spreadsheet formula evaluation in fields imported from data files."""

        text = str(value or "")
        return f"'{text}" if text.lstrip().startswith(("=", "+", "-", "@")) else text

    @staticmethod
    def _unique_sources(sources: Iterable[Mapping[str, Any]]) -> list[dict[str, Any]]:
        seen: set[str] = set()
        output: list[dict[str, Any]] = []
        for source in sources:
            identity = str(source.get("url") or source.get("source_id") or "")
            if not identity or identity in seen:
                continue
            seen.add(identity)
            output.append(dict(source))
        return output
